"""Per-account bank-statement workbooks.

The weekly report (``report.py``) is an *analytical* roll-up of all accounts in
one workbook. This module is the complementary *statement* view: one Excel file
**per account**, each a plain line-by-line transaction listing with a running
balance — exactly like the statement your bank emails you.

Reads the build-layer ``transactions_flow`` view (never the raw table), so the
statement uses the same ``effective_date`` and internal-transfer classification
as every other consumer. The aggregation logic is pure (DataFrame in →
statement DataFrame out) so it is unit tested without a database, mirroring
``report.py``.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
import psycopg
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import MONEY_FORMAT, db
from .config import Settings

log = logging.getLogger(__name__)

# Order matters: this is the on-statement column layout.
STATEMENT_COLUMNS = ["Date", "Description", "Category", "Debit", "Credit", "Balance"]
_MONEY_COLS = ("Debit", "Credit", "Balance")


@dataclass(frozen=True)
class Account:
    account_id: str
    account_number: str
    account_name: str
    currency: str = "ZAR"


def load_accounts(conn: psycopg.Connection) -> list[Account]:
    """List every known account, with its latest snapshot currency."""
    query = """
        select a.account_id,
               coalesce(a.account_number, a.account_id) as account_number,
               coalesce(a.account_name, '')             as account_name,
               coalesce(
                   (select b.currency from balances b
                     where b.account_id = a.account_id
                     order by b.as_of_date desc, b.captured_at desc
                     limit 1),
                   'ZAR'
               )                                        as currency
        from accounts a
        order by account_number;
    """
    with conn.cursor() as cur:
        cur.execute(query)
        rows = cur.fetchall()
    return [Account(*row) for row in rows]


def load_account_transactions(
    conn: psycopg.Connection, account_id: str, start: date | None, end: date
) -> pd.DataFrame:
    """Read one account's transactions up to ``end``, in statement order.

    ``start`` bounds the period below; pass ``None`` for the full history (from
    the account's first transaction). Reads ``transactions_flow`` so each row
    carries the shared ``effective_date`` and ``flow_type``; joins the base
    table only for ``day_seq``/``ingested_at`` so same-day postings keep the
    API's stable order — which is also the order the running balance was
    computed in. (Display order, newest-first, is applied later in
    :func:`build_account_statement`.)
    """
    params: list = [account_id, end]
    lower_bound = ""
    if start is not None:
        lower_bound = "and f.effective_date >= %s"
        params.append(start)
    query = f"""
        select f.effective_date, f.description, f.transaction_type, f.type,
               f.amount, f.running_balance, f.category, f.flow_type, t.day_seq
        from transactions_flow f
        join transactions t using (transaction_hash)
        where f.account_id = %s and f.effective_date <= %s {lower_bound}
        order by f.effective_date, t.day_seq, t.ingested_at;
    """
    cols = ["effective_date", "description", "transaction_type", "type",
            "amount", "running_balance", "category", "flow_type", "day_seq"]
    with conn.cursor() as cur:
        cur.execute(query, tuple(params))
        rows = cur.fetchall()
    df = pd.DataFrame(rows, columns=cols)
    if not df.empty:
        df["amount"] = df["amount"].astype(float)
        df["running_balance"] = df["running_balance"].astype(float)
        df["effective_date"] = pd.to_datetime(df["effective_date"])
    return df


def load_opening_balance(
    conn: psycopg.Connection, account_id: str, start: date
) -> float | None:
    """The running balance on the last transaction *before* the period.

    This is the statement's opening balance — the carried-forward figure a bank
    prints above the first line. ``None`` when there is no prior balance to
    anchor to (the period contains the account's earliest data).
    """
    query = """
        select f.running_balance
        from transactions_flow f
        join transactions t using (transaction_hash)
        where f.account_id = %s and f.effective_date < %s
          and f.running_balance is not null
        order by f.effective_date desc, t.day_seq desc, t.ingested_at desc
        limit 1;
    """
    with conn.cursor() as cur:
        cur.execute(query, (account_id, start))
        row = cur.fetchone()
    return float(row[0]) if row and row[0] is not None else None


def build_account_statement(
    df: pd.DataFrame,
    opening_balance: float | None = None,
    newest_first: bool = True,
) -> dict:
    """Turn one account's transactions into a bank-statement DataFrame + totals.

    Debits (outflows) and credits (inflows) are split into their own columns and
    a running ``Balance`` is carried down each row. The bank's own
    ``running_balance`` is authoritative and used as-is when present; gaps are
    filled arithmetically (previous balance + signed amount). When there is no
    anchor at all — no opening balance and no bank balance yet — the Balance
    cell is left blank rather than fabricated from zero, until the first
    authoritative balance appears.

    Internal transfers (``flow_type == 'internal_transfer'``, from the
    ``transactions_flow`` view) are reported separately from real money in/out,
    matching the weekly report. A frame without ``flow_type`` is tolerated by
    treating every row as external.

    The running balance is always computed forward (oldest → newest); when
    ``newest_first`` is true the finished rows are then reversed for display so
    the most recent transaction sits at the top, the way a bank statement reads.
    Each row keeps its own balance regardless of display order.

    Pure (no I/O) so it is unit tested without a database.
    """
    if df.empty:
        return {
            "statement": pd.DataFrame(columns=STATEMENT_COLUMNS),
            "opening_balance": opening_balance,
            "closing_balance": opening_balance,
            "money_in": 0.0,
            "money_out": 0.0,
            "transfers_in": 0.0,
            "transfers_out": 0.0,
            "count": 0,
        }

    sort_cols = [c for c in ("effective_date", "day_seq") if c in df.columns]
    df = df.sort_values(sort_cols).reset_index(drop=True)

    balances: list[float | None] = []
    prev = opening_balance
    for amount, running in zip(df["amount"], df["running_balance"]):
        if running is not None and not pd.isna(running):
            balance = round(float(running), 2)
        elif prev is not None:
            balance = round(prev + float(amount), 2)
        else:
            balance = None  # no anchor yet — leave blank, don't fabricate from 0
        balances.append(balance)
        if balance is not None:
            prev = balance
    if balances[0] is None:
        log.warning(
            "Statement starts before the first bank-reported balance; "
            "%d leading Balance cell(s) left blank",
            sum(1 for b in balances if b is None),
        )

    statement = pd.DataFrame(
        {
            "Date": df["effective_date"].dt.date,
            "Description": df["description"].fillna("").astype(str).str.strip(),
            "Category": df.get("category", pd.Series([""] * len(df))).fillna(""),
            "Debit": [round(-a, 2) if a < 0 else None for a in df["amount"]],
            "Credit": [round(a, 2) if a > 0 else None for a in df["amount"]],
            "Balance": balances,
        }
    )

    # If we were not handed an opening balance, infer it by stepping back off the
    # first row so the header reconciles with the first printed line — but only
    # when that first line actually has a balance to step back from.
    if opening_balance is None and balances[0] is not None:
        opening_balance = round(balances[0] - float(df["amount"].iloc[0]), 2)

    # Internal transfers net to zero across your own accounts and would otherwise
    # inflate both money-in and money-out, so they get their own totals.
    if "flow_type" in df.columns:
        is_transfer = df["flow_type"].eq("internal_transfer")
    else:
        is_transfer = pd.Series([False] * len(df))
    outflow = df["amount"] < 0
    inflow = df["amount"] > 0
    money_out = round(float(df.loc[outflow & ~is_transfer, "amount"].abs().sum()), 2)
    money_in = round(float(df.loc[inflow & ~is_transfer, "amount"].sum()), 2)
    transfers_out = round(float(df.loc[outflow & is_transfer, "amount"].abs().sum()), 2)
    transfers_in = round(float(df.loc[inflow & is_transfer, "amount"].sum()), 2)
    closing_balance = balances[-1]  # newest balance, before any display reorder

    # Newest at the top for display; balances were computed forward above so each
    # row already carries the correct figure.
    if newest_first:
        statement = statement.iloc[::-1].reset_index(drop=True)

    return {
        "statement": statement,
        "opening_balance": opening_balance,
        "closing_balance": closing_balance,
        "money_in": money_in,
        "money_out": money_out,
        "transfers_in": transfers_in,
        "transfers_out": transfers_out,
        "count": len(statement),
    }


def _safe_filename(text: str) -> str:
    """Make a string safe for use as a filename component."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", text.strip())
    return cleaned.strip("-") or "account"


def write_statement_workbook(
    statement: dict, account: Account, start: date, end: date, path: Path
) -> Path:
    """Write a single-account, bank-statement-style .xlsx.

    Header block (account, period, opening/closing balance, totals) sits above
    the line-by-line transaction table, the way a printed statement reads.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    frame: pd.DataFrame = statement["statement"]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        sheet_name = "Statement"
        header_rows = 11  # account/period/balances block before the table
        body = frame if not frame.empty else pd.DataFrame(
            [{"Date": None, "Description": "No transactions for this period",
              "Category": None, "Debit": None, "Credit": None, "Balance": None}]
        )
        body = body.reindex(columns=STATEMENT_COLUMNS)
        body.to_excel(writer, sheet_name=sheet_name, index=False, startrow=header_rows)
        ws = writer.sheets[sheet_name]

        ccy = account.currency or "ZAR"
        meta = [
            ("Account statement", ""),
            ("Account name", account.account_name),
            ("Account number", account.account_number),
            ("Period", f"{start.isoformat()} to {end.isoformat()}"),
            ("Currency", ccy),
            ("Opening balance", statement["opening_balance"]),
            ("Money in (excl. transfers)", statement["money_in"]),
            ("Money out (excl. transfers)", statement["money_out"]),
            ("Transfers in", statement["transfers_in"]),
            ("Transfers out", statement["transfers_out"]),
            ("Closing balance", statement["closing_balance"]),
        ]
        for i, (label, value) in enumerate(meta, start=1):
            label_cell = ws.cell(row=i, column=1, value=label)
            value_cell = ws.cell(row=i, column=2, value=value)
            label_cell.font = Font(bold=True)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                value_cell.number_format = MONEY_FORMAT
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)

        # Style the transaction-table header row (1-based; +1 for the header itself).
        head_row = header_rows + 1
        fill = PatternFill("solid", fgColor="DDEBF7")
        for col_idx, _name in enumerate(STATEMENT_COLUMNS, start=1):
            cell = ws.cell(row=head_row, column=col_idx)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

        # Money formatting + sensible column widths.
        money_cols = {STATEMENT_COLUMNS.index(c) + 1 for c in _MONEY_COLS}
        first_data_row = head_row + 1
        last_data_row = head_row + max(len(body), 1)
        date_col = STATEMENT_COLUMNS.index("Date") + 1
        widths = {1: 12, 2: 42, 3: 16, 4: 14, 5: 14, 6: 16}
        for col_idx in range(1, len(STATEMENT_COLUMNS) + 1):
            letter = get_column_letter(col_idx)
            fmt = MONEY_FORMAT if col_idx in money_cols else (
                "yyyy-mm-dd" if col_idx == date_col else None
            )
            if fmt:
                for r in range(first_data_row, last_data_row + 1):
                    ws.cell(row=r, column=col_idx).number_format = fmt
            ws.column_dimensions[letter].width = widths.get(col_idx, 14)

        ws.freeze_panes = ws.cell(row=first_data_row, column=1)

    return path


def generate_account_statements(
    settings: Settings,
    end: date | None = None,
    out_dir: Path | None = None,
    days: int | None = None,
    include_patterns: list[str] | None = None,
    exclude_patterns: list[str] | None = None,
) -> tuple[list[Path], dict]:
    """Build one bank-statement workbook per account up to ``end``.

    By default (``days`` is ``None``) each statement spans the account's **full
    history** — its first stored transaction through ``end`` — so the weekly run
    regenerates a complete, up-to-date statement each time. Pass ``days`` to
    limit it to a trailing window instead. Rows are newest-first.

    ``include_patterns`` / ``exclude_patterns`` are case-insensitive substrings
    matched against ``account_name``; exclude is applied first.  Pass neither to
    include all accounts (original behaviour).

    Returns ``(paths, info)``. Accounts with no transactions are skipped so the
    email only carries statements that have activity.
    """
    end = end or date.today()
    start = None if days is None else end - timedelta(days=days - 1)
    out_dir = out_dir or Path("reports")

    # One pooled connection for the whole pass (accounts + per-account reads)
    # instead of a fresh connect per query — kinder to the transaction pooler.
    loaded: list[tuple[Account, pd.DataFrame, float | None]] = []
    with db.connect(settings.reporting_db_url) as conn:
        for account in load_accounts(conn):
            if exclude_patterns and any(
                p.lower() in account.account_name.lower() for p in exclude_patterns
            ):
                log.info("Excluding account %s from statements", account.account_number)
                continue
            if include_patterns and not any(
                p.lower() in account.account_name.lower() for p in include_patterns
            ):
                continue
            df = load_account_transactions(conn, account.account_id, start, end)
            if df.empty:
                log.info("No transactions for %s; skipping statement",
                         account.account_number)
                continue
            # Full-history statements start at the account's first transaction,
            # so there is nothing before the period to anchor an opening to.
            opening = (
                load_opening_balance(conn, account.account_id, start)
                if start is not None else None
            )
            loaded.append((account, df, opening))

    paths: list[Path] = []
    per_account: list[dict] = []
    earliest: date | None = None
    for account, df, opening in loaded:
        period_start = start if start is not None else df["effective_date"].min().date()
        earliest = period_start if earliest is None else min(earliest, period_start)
        statement = build_account_statement(df, opening)
        filename = (
            f"statement_{_safe_filename(account.account_number)}_"
            f"{period_start.isoformat()}_to_{end.isoformat()}.xlsx"
        )
        path = write_statement_workbook(
            statement, account, period_start, end, out_dir / filename
        )
        paths.append(path)
        per_account.append(
            {
                "account_number": account.account_number,
                "account_name": account.account_name,
                "start": period_start.isoformat(),
                "transactions": statement["count"],
                "closing_balance": statement["closing_balance"],
            }
        )
        log.info("Wrote statement %s (%d transactions)", path, statement["count"])

    info = {
        "start": earliest.isoformat() if earliest else end.isoformat(),
        "end": end.isoformat(),
        "accounts": len(paths),
        "full_history": days is None,
        "per_account": per_account,
    }
    return paths, info
