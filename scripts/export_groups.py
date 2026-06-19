"""Generate an Excel workbook showing all accounts and their current group assignments.

Usage (from repo root, with DATABASE_URL set):
    python scripts/export_groups.py [output.xlsx]

The output file has two sheets:
  "Accounts"  — one row per account with current group and a "New Group" column to edit
  "Groups"    — the current group definitions for reference

After editing, send the file back so groups.py can be updated.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

# Make sure invespend package is importable when run from repo root.
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import psycopg
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from invespend.groups import EXCLUDE_ACCOUNTS, GROUPS


def _assign_group(account_name: str, account_number: str) -> str:
    name_lower = account_name.lower()
    for excl in EXCLUDE_ACCOUNTS:
        if excl.lower() in name_lower:
            return "EXCLUDE"
    for g in GROUPS:
        if g.matches(account_name, account_number):
            return g.name
    return "default"


def _header_fill(hex_colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_colour)


def _thin_border() -> Border:
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def build_workbook(rows: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Accounts ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Accounts"

    group_names = ["default"] + [g.name for g in GROUPS] + ["EXCLUDE"]

    headers = [
        "account_number",
        "account_name",
        "product_name",
        "reference_name",
        "current_group",
        "new_group  ← EDIT THIS",
        "new_recipients  ← EDIT THIS (comma-separated emails)",
    ]
    header_fill = _header_fill("2E75B6")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _thin_border()

    # Group colour map
    palette = {
        "default":  "D9E1F2",
        "EXCLUDE":  "FFCCCC",
    }
    group_colours = [
        "E2EFDA", "FFF2CC", "FCE4D6", "EAD1DC", "D9D2E9",
    ]
    for i, g in enumerate(GROUPS):
        palette[g.name] = group_colours[i % len(group_colours)]

    # Current recipients per group
    recipients_map = {g.name: ", ".join(g.recipients) for g in GROUPS}

    for row_idx, acc in enumerate(rows, 2):
        grp = _assign_group(acc["account_name"] or "", acc["account_number"] or "")
        fill = PatternFill("solid", fgColor=palette.get(grp, "FFFFFF"))
        border = _thin_border()

        values = [
            acc["account_number"],
            acc["account_name"],
            acc["product_name"],
            acc["reference_name"],
            grp,
            grp,                                         # pre-fill new_group with current
            recipients_map.get(grp, ""),                 # pre-fill recipients
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(wrap_text=False)

    # Column widths
    widths = [22, 40, 25, 30, 18, 22, 50]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # Legend
    legend_row = len(rows) + 3
    ws.cell(row=legend_row, column=1, value="LEGEND").font = Font(bold=True)
    for i, (grp_name, colour) in enumerate(palette.items(), 1):
        cell = ws.cell(row=legend_row + i, column=1, value=grp_name)
        cell.fill = PatternFill("solid", fgColor=colour)
        cell.border = _thin_border()

    ws.cell(row=legend_row, column=3,
            value="Valid values for 'new_group': " + ", ".join(group_names)).font = Font(italic=True)

    # ── Sheet 2: Groups ────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Groups")
    g_headers = ["group_name", "recipients", "name_patterns", "account_numbers"]
    for col, h in enumerate(g_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = _thin_border()

    for row_idx, g in enumerate(GROUPS, 2):
        ws2.cell(row=row_idx, column=1, value=g.name)
        ws2.cell(row=row_idx, column=2, value=", ".join(g.recipients))
        ws2.cell(row=row_idx, column=3, value=", ".join(g.name_patterns))
        ws2.cell(row=row_idx, column=4, value=", ".join(g.account_numbers))
        for col in range(1, 5):
            ws2.cell(row=row_idx, column=col).border = _thin_border()

    # Excluded patterns row
    excl_row = len(GROUPS) + 3
    ws2.cell(row=excl_row, column=1, value="EXCLUDE_ACCOUNTS").font = Font(bold=True)
    ws2.cell(row=excl_row, column=2, value=", ".join(EXCLUDE_ACCOUNTS))

    for col, w in enumerate([20, 45, 50, 40], 1):
        ws2.column_dimensions[get_column_letter(col)].width = w

    return wb


def main() -> None:
    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("groups_export.xlsx")

    database_url = os.environ.get("DATABASE_URL") or os.environ.get("REPORT_DATABASE_URL")
    if not database_url:
        raise SystemExit("DATABASE_URL is not set")

    with psycopg.connect(database_url, prepare_threshold=None) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "select account_number, account_name, product_name, reference_name "
                "from accounts "
                "order by account_name;"
            )
            cols = [d[0] for d in cur.description]
            rows = [dict(zip(cols, row)) for row in cur.fetchall()]

    wb = build_workbook(rows)
    wb.save(out_path)
    print(f"Saved {len(rows)} accounts → {out_path}")


if __name__ == "__main__":
    main()
