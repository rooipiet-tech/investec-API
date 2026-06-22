import datetime as dt

import pandas as pd
from openpyxl import load_workbook

from invespend import MONEY_FORMAT
from invespend.statements import (
    Account,
    build_account_statement,
    write_statement_workbook,
    _safe_filename,
)


def _sample_df():
    # Same-day pair carries running_balance; one row has a NULL balance to
    # exercise the arithmetic gap-fill.
    return pd.DataFrame(
        {
            "effective_date": pd.to_datetime(
                ["2026-05-25", "2026-05-26", "2026-05-26", "2026-05-27"]
            ),
            "description": ["Woolworths", "Uber", "Salary", "Netflix"],
            "transaction_type": ["CardPurchases"] * 4,
            "type": ["DEBIT", "DEBIT", "CREDIT", "DEBIT"],
            "amount": [-200.0, -85.5, 15000.0, -199.0],
            "running_balance": [800.0, 714.5, 15714.5, None],
            "category": ["Groceries", "Transport", "Income", "Subscriptions"],
            "flow_type": ["external_outflow", "external_outflow",
                          "external_inflow", "external_outflow"],
            "day_seq": [0, 0, 1, 0],
        }
    )


def _sample_df_with_transfers():
    # Adds an internal transfer in and out alongside real spend/income, flagged
    # the way the transactions_flow view flags them.
    return pd.DataFrame(
        {
            "effective_date": pd.to_datetime(
                ["2026-05-25", "2026-05-26", "2026-05-26", "2026-05-27"]
            ),
            "description": ["Woolworths", "Transfer to savings", "Salary",
                            "Transfer from savings"],
            "transaction_type": ["CardPurchases", "Transfer", "Salary", "Transfer"],
            "type": ["DEBIT", "DEBIT", "CREDIT", "CREDIT"],
            "amount": [-200.0, -1000.0, 15000.0, 500.0],
            "running_balance": [800.0, -200.0, 14800.0, 15300.0],
            "category": ["Groceries", "Transfers", "Income", "Transfers"],
            "flow_type": ["external_outflow", "internal_transfer",
                          "external_inflow", "internal_transfer"],
            "day_seq": [0, 0, 1, 0],
        }
    )


def test_newest_first_ordering():
    # Default display is newest-first: the latest posting sits on row 0, the
    # earliest at the bottom.
    stmt = build_account_statement(_sample_df(), opening_balance=1000.0)["statement"]
    assert stmt.loc[0, "Description"] == "Netflix"      # 2026-05-27, newest
    assert stmt.loc[3, "Description"] == "Woolworths"   # 2026-05-25, oldest


def test_debit_credit_split():
    stmt = build_account_statement(_sample_df(), opening_balance=1000.0)["statement"]
    # Debits populate Debit (positive), leave Credit empty, and vice versa.
    # (Rows are newest-first, so Woolworths — the oldest — is the last row.)
    assert stmt.loc[3, "Debit"] == 200.0 and pd.isna(stmt.loc[3, "Credit"])
    assert stmt.loc[1, "Credit"] == 15000.0 and pd.isna(stmt.loc[1, "Debit"])


def test_uses_bank_running_balance_and_fills_gaps():
    result = build_account_statement(_sample_df(), opening_balance=1000.0)
    stmt = result["statement"]
    # Balances are computed forward then displayed newest-first. The NULL final
    # balance is filled (15714.5 + -199.0 = 15515.5) and lands on the top row;
    # the stored balances follow below in reverse chronological order.
    assert list(stmt["Balance"]) == [15515.5, 15714.5, 714.5, 800.0]
    assert result["closing_balance"] == 15515.5


def test_oldest_first_when_disabled():
    stmt = build_account_statement(
        _sample_df(), opening_balance=1000.0, newest_first=False
    )["statement"]
    assert list(stmt["Balance"]) == [800.0, 714.5, 15714.5, 15515.5]
    assert stmt.loc[0, "Description"] == "Woolworths"


def test_totals_and_opening_passthrough():
    result = build_account_statement(_sample_df(), opening_balance=1000.0)
    assert result["opening_balance"] == 1000.0
    # No transfers in this sample, so money in/out equal the raw debit/credit sums.
    assert result["money_out"] == 484.5      # 200 + 85.5 + 199
    assert result["money_in"] == 15000.0
    assert result["transfers_out"] == 0.0
    assert result["transfers_in"] == 0.0
    assert result["count"] == 4


def test_internal_transfers_separated_from_money():
    result = build_account_statement(_sample_df_with_transfers(), opening_balance=1000.0)
    # flow_type == internal_transfer rows are pulled out of money in/out and
    # reported on their own, matching the weekly report's classification.
    assert result["money_out"] == 200.0        # Woolworths only (transfer excluded)
    assert result["money_in"] == 15000.0       # Salary only (transfer excluded)
    assert result["transfers_out"] == 1000.0
    assert result["transfers_in"] == 500.0


def test_frame_without_flow_type_treated_as_external():
    df = _sample_df().drop(columns=["flow_type"])
    result = build_account_statement(df, opening_balance=1000.0)
    assert result["money_out"] == 484.5
    assert result["transfers_out"] == 0.0


def test_opening_balance_inferred_when_absent():
    # No opening handed in: infer it by stepping back off the first row.
    result = build_account_statement(_sample_df(), opening_balance=None)
    # First stored balance 800 minus the first amount (-200) => 1000.
    assert result["opening_balance"] == 1000.0


def test_unanchored_balances_left_blank_not_zeroed():
    # No opening balance AND the first rows carry no bank balance: the Balance
    # column must stay blank (not fabricated from 0) until the first
    # authoritative running_balance appears, then carry forward from it.
    df = _sample_df()
    df["running_balance"] = [None, None, 15714.5, None]
    result = build_account_statement(df, opening_balance=None, newest_first=False)
    stmt = result["statement"]
    assert pd.isna(stmt.loc[0, "Balance"])  # blank cell in the workbook
    assert pd.isna(stmt.loc[1, "Balance"])
    assert list(stmt["Balance"])[2:] == [15714.5, 15515.5]
    assert result["opening_balance"] is None
    assert result["closing_balance"] == 15515.5


def test_empty_df_is_safe():
    result = build_account_statement(pd.DataFrame())
    assert result["count"] == 0
    assert result["money_out"] == 0.0
    assert result["transfers_in"] == 0.0
    assert list(result["statement"].columns) == [
        "Date", "Description", "Category", "Debit", "Credit", "Balance",
    ]


def test_write_statement_workbook(tmp_path):
    result = build_account_statement(_sample_df(), opening_balance=1000.0)
    account = Account("acc-1", "10010000001", "Primary Cheque", "ZAR")
    out = write_statement_workbook(
        result, account, dt.date(2026, 5, 25), dt.date(2026, 5, 31),
        tmp_path / "stmt.xlsx",
    )
    assert out.exists() and out.stat().st_size > 0

    wb = load_workbook(out)
    ws = wb["Statement"]
    # Header block (14 rows: 11 account/period/balances + 3 reconciliation).
    header = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(1, 15)}
    assert header["Account number"] == "10010000001"
    assert header["Closing balance"] == 15515.5
    assert header["Transfers in"] == 0.0
    # No balance snapshot → reconciliation row shows N/A.
    assert header["Reconciled"] == "N/A — no balance snapshot"
    # The transaction table header is present below the 14-row metadata block.
    assert ws.cell(row=15, column=1).value == "Date"
    assert ws.cell(row=15, column=6).value == "Balance"
    # Money cells (header block + Balance column) carry the Accounting format.
    closing_row = next(r for r in range(1, 15)
                       if ws.cell(row=r, column=1).value == "Closing balance")
    assert ws.cell(row=closing_row, column=2).number_format == MONEY_FORMAT
    assert ws.cell(row=16, column=6).number_format == MONEY_FORMAT  # first Balance cell


def test_write_statement_workbook_empty(tmp_path):
    result = build_account_statement(pd.DataFrame())
    account = Account("acc-1", "10010000001", "Primary Cheque", "ZAR")
    out = write_statement_workbook(
        result, account, dt.date(2026, 5, 25), dt.date(2026, 5, 31),
        tmp_path / "empty.xlsx",
    )
    ws = load_workbook(out)["Statement"]
    assert ws.cell(row=16, column=2).value == "No transactions for this period"


def _recon_result(reconciled, recon_difference, investec_balance=None, investec_date=None):
    result = build_account_statement(_sample_df(), opening_balance=1000.0)
    result["reconciled"] = reconciled
    result["recon_difference"] = recon_difference
    result["investec_balance"] = investec_balance
    result["investec_balance_date"] = investec_date
    return result


def test_write_statement_workbook_recon_pass(tmp_path):
    result = _recon_result(True, 0.0, investec_balance=15515.5,
                           investec_date=dt.date(2026, 5, 31))
    account = Account("acc-1", "10010000001", "Primary Cheque", "ZAR")
    out = write_statement_workbook(
        result, account, dt.date(2026, 5, 25), dt.date(2026, 5, 31),
        tmp_path / "recon_pass.xlsx",
    )
    ws = load_workbook(out)["Statement"]
    assert ws.cell(row=14, column=2).value == "RECONCILED"
    # Green fill applied
    assert ws.cell(row=14, column=2).fill.fgColor.rgb == "FF" + "C6EFCE"


def test_write_statement_workbook_recon_fail(tmp_path):
    result = _recon_result(False, -150.0, investec_balance=15665.5,
                           investec_date=dt.date(2026, 5, 31))
    account = Account("acc-1", "10010000001", "Primary Cheque", "ZAR")
    out = write_statement_workbook(
        result, account, dt.date(2026, 5, 25), dt.date(2026, 5, 31),
        tmp_path / "recon_fail.xlsx",
    )
    ws = load_workbook(out)["Statement"]
    recon_cell = ws.cell(row=14, column=2)
    assert "DIFFERENCE" in recon_cell.value
    assert "-150" in recon_cell.value
    # Red fill applied
    assert recon_cell.fill.fgColor.rgb == "FF" + "FFC7CE"


def test_write_statement_workbook_recon_none(tmp_path):
    result = _recon_result(None, None)
    account = Account("acc-1", "10010000001", "Primary Cheque", "ZAR")
    out = write_statement_workbook(
        result, account, dt.date(2026, 5, 25), dt.date(2026, 5, 31),
        tmp_path / "recon_none.xlsx",
    )
    ws = load_workbook(out)["Statement"]
    assert ws.cell(row=14, column=2).value == "N/A — no balance snapshot"


def test_intraday_chain_sort():
    # Three same-day transactions whose day_seq is WRONG (reversed from Investec
    # order) but whose running_balance encodes the correct sequence.
    # Correct sequence: Woolworths(800) → Uber(714.5) → Salary(15714.5).
    # Wrong day_seq order puts Salary first (seq 0) and Woolworths last (seq 2).
    df = pd.DataFrame(
        {
            "effective_date": pd.to_datetime(["2026-05-26"] * 3),
            "description": ["Salary", "Uber", "Woolworths"],
            "transaction_type": ["Salary", "CardPurchases", "CardPurchases"],
            "type": ["CREDIT", "DEBIT", "DEBIT"],
            "amount": [15000.0, -85.5, -200.0],
            "running_balance": [15714.5, 714.5, 800.0],
            "category": ["Income", "Transport", "Groceries"],
            "flow_type": ["external_inflow", "external_outflow", "external_outflow"],
            "day_seq": [0, 1, 2],  # intentionally wrong
        }
    )
    result = build_account_statement(df, opening_balance=1000.0, newest_first=False)
    stmt = result["statement"]
    # Chain sort must restore Investec's sequence regardless of day_seq.
    assert list(stmt["Description"]) == ["Woolworths", "Uber", "Salary"]
    assert list(stmt["Balance"]) == [800.0, 714.5, 15714.5]


def test_intraday_chain_sort_with_null_rb():
    # Two same-day transactions have running_balance; one has NULL.
    # Wrong day_seq order: Netflix(seq=0), Woolworths(seq=1), Uber(seq=2).
    # Correct Investec sequence: Woolworths(800) → Uber(714.5) → Netflix(NULL fills to 515.5).
    # Arithmetic check: Netflix(-199) at gap 0: 1000-199=801 ≠ Woolworths.before=1000. No match.
    #                   Netflix at gap 1: 800-199=601 ≠ Uber.before=800. No match.
    #                   No unique match → append after chain → [Woolworths, Uber, Netflix].
    df = pd.DataFrame(
        {
            "effective_date": pd.to_datetime(["2026-05-26"] * 3),
            "description": ["Netflix", "Woolworths", "Uber"],
            "transaction_type": ["Subscription", "CardPurchases", "CardPurchases"],
            "type": ["DEBIT", "DEBIT", "DEBIT"],
            "amount": [-199.0, -200.0, -85.5],
            "running_balance": [None, 800.0, 714.5],
            "category": ["Subscriptions", "Groceries", "Transport"],
            "flow_type": ["external_outflow", "external_outflow", "external_outflow"],
            "day_seq": [0, 1, 2],  # wrong order
        }
    )
    result = build_account_statement(df, opening_balance=1000.0, newest_first=False)
    stmt = result["statement"]
    descriptions = list(stmt["Description"])
    # Arithmetic placement: no gap is consistent → Netflix appended after the chain.
    assert list(descriptions) == ["Woolworths", "Uber", "Netflix"]
    assert list(stmt["Balance"]) == [800.0, 714.5, 515.5]


def test_intraday_chain_sort_no_anchor():
    # Three same-day transactions with no opening_balance (full-history statement).
    # All have running_balance, so orphan detection can find the chain start.
    # Correct sequence: Woolworths(800) → Uber(714.5) → Salary(15714.5).
    # Wrong day_seq order: Salary(seq=0), Uber(seq=1), Woolworths(seq=2).
    df = pd.DataFrame(
        {
            "effective_date": pd.to_datetime(["2026-05-26"] * 3),
            "description": ["Salary", "Uber", "Woolworths"],
            "transaction_type": ["Salary", "CardPurchases", "CardPurchases"],
            "type": ["CREDIT", "DEBIT", "DEBIT"],
            "amount": [15000.0, -85.5, -200.0],
            "running_balance": [15714.5, 714.5, 800.0],
            "category": ["Income", "Transport", "Groceries"],
            "flow_type": ["external_inflow", "external_outflow", "external_outflow"],
            "day_seq": [0, 1, 2],  # wrong
        }
    )
    result = build_account_statement(df, opening_balance=None, newest_first=False)
    stmt = result["statement"]
    # Orphan: 800-(-200)=1000 not in {15714.5, 714.5, 800} → Woolworths is orphan → first.
    # Then 800→Uber(800-85.5=714.5)→Salary(714.5+15000=15714.5).
    assert list(stmt["Description"]) == ["Woolworths", "Uber", "Salary"]
    assert list(stmt["Balance"]) == [800.0, 714.5, 15714.5]


def test_intraday_chain_sort_no_anchor_with_null_rb():
    # opening=None (full-history) AND one NULL-rb row — the previously unhandled case.
    # Correct Investec order: Woolworths(800) → Uber(714.5) → Netflix(NULL fills to 515.5).
    # Wrong day_seq order: Uber(0), Woolworths(1), Netflix(2).
    # Orphan detection on has_rb=[Uber, Woolworths]:
    #   Woolworths before = 800−(−200) = 1000, not in {800, 714.5} → orphan → anchor=1000.
    #   Uber before = 714.5−(−85.5) = 800, IS in {800, 714.5} → not orphan.
    # Chain: Woolworths(800) → Uber(714.5).
    # Netflix day_seq=2 > chained_seqs [1, 0] → appended last.
    # Result order: Woolworths, Uber, Netflix. Netflix fill = 714.5 + (−199) = 515.5.
    df = pd.DataFrame(
        {
            "effective_date": pd.to_datetime(["2026-05-26"] * 3),
            "description": ["Uber", "Woolworths", "Netflix"],
            "transaction_type": ["CardPurchases", "CardPurchases", "Subscription"],
            "type": ["DEBIT", "DEBIT", "DEBIT"],
            "amount": [-85.5, -200.0, -199.0],
            "running_balance": [714.5, 800.0, None],
            "category": ["Transport", "Groceries", "Subscriptions"],
            "flow_type": ["external_outflow", "external_outflow", "external_outflow"],
            "day_seq": [0, 1, 2],  # wrong order — Uber(0) should come after Woolworths(1)
        }
    )
    result = build_account_statement(df, opening_balance=None, newest_first=False)
    stmt = result["statement"]
    descriptions = list(stmt["Description"])
    assert descriptions.index("Woolworths") < descriptions.index("Uber"), descriptions
    assert descriptions.index("Uber") < descriptions.index("Netflix"), descriptions
    wool_idx = descriptions.index("Woolworths")
    uber_idx = descriptions.index("Uber")
    netflix_idx = descriptions.index("Netflix")
    assert stmt.loc[wool_idx, "Balance"] == 800.0
    assert stmt.loc[uber_idx, "Balance"] == 714.5
    assert stmt.loc[netflix_idx, "Balance"] == 515.5


def test_safe_filename():
    assert _safe_filename("1001 0000/001") == "1001-0000-001"
    assert _safe_filename("   ") == "account"
