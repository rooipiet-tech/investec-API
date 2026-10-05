"""Beneficiary matching must not change ingest, report or statement outputs.

F9: ingest stdout / summary / sync_runs identical with the feature OFF and ON
(including when the beneficiary fetch fails). F10: Excel outputs identical to a
golden digest generated from the unmodified base code. F11: OFF by default.
No live DB or API: everything is faked. Synthetic data only.
"""
import argparse
import datetime as dt
import json
import logging
from contextlib import contextmanager
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from invespend import cli, db, ingest
from invespend.config import Settings

FIXTURES = Path(__file__).parent / "fixtures" / "beneficiary"
SUMMARY_KEYS = ["from_date", "to_date", "accounts", "transactions_upserted",
                "balances_captured", "new_accounts"]

ACCOUNTS = [{"accountId": "acc-1", "accountNumber": "TEST-ACCT-01",
             "accountName": "Test Cheque"}]
TXS = [
    {"type": "DEBIT", "transactionType": "OnlineBankingPayments",
     "description": "PAYMENT TO ACME TRADING", "amount": 100.0,
     "postingDate": "2026-01-02", "valueDate": "2026-01-02", "actionDate": "2026-01-02"},
]
BENES = [{"beneficiaryId": "ben-acme", "beneficiaryName": "Acme Trading",
          "accountNumber": "9999 0000 1234"}]


def _settings(enabled: bool) -> Settings:
    return Settings(
        investec_client_id="test-id", investec_client_secret="test-secret",
        investec_api_key="test-key", investec_base_url="https://example.invalid",
        database_url="postgresql://fake", smtp_host="smtp.example.invalid",
        smtp_port=587, smtp_user="", smtp_password="", report_sender="",
        beneficiary_matching_enabled=enabled,
    )


class _Cursor:
    def __init__(self, log):
        self._log = log
        self._rows = []
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def execute(self, sql, params=None):
        self._log.append(sql)
        low = " ".join(sql.lower().split())
        self._rows = []
        if low.startswith("select account_number from accounts"):
            self._rows = [("TEST-ACCT-01",)]
        elif low.startswith("select beneficiary_id,"):
            self._rows = [("ben-acme", "Acme Trading", None, None, None, None,
                           "9999 0000 1234")]
        elif "from transactions_flow" in low:
            self._rows = [("h1", "DEBIT", "OnlineBankingPayments",
                           "PAYMENT TO ACME TRADING", -100.0, "external_outflow")]
        self.rowcount = 1

    def fetchall(self):
        return self._rows


class _Conn:
    def __init__(self, log):
        self._log = log

    def cursor(self):
        return _Cursor(self._log)


class _World:
    """Records everything the ingest run does against the fakes."""

    def __init__(self, bene_data=None, bene_raises=None):
        self.events = []
        self.sql = []
        self.finish_calls = []
        self.bene_calls = 0
        self.bene_data = BENES if bene_data is None else bene_data
        self.bene_raises = bene_raises


@pytest.fixture
def world_factory(monkeypatch):
    def make(enabled, bene_data=None, bene_raises=None):
        w = _World(bene_data, bene_raises)

        class _FakeClient:
            def __init__(self, **_kw):
                pass

            def get_accounts(self):
                return ACCOUNTS

            def get_balance(self, _account_id):
                return {"currentBalance": 1.0}

            def get_transactions(self, _account_id, _start, _end):
                return TXS

            def get_beneficiaries(self):
                w.bene_calls += 1
                w.events.append("get_beneficiaries")
                if w.bene_raises:
                    raise w.bene_raises
                return w.bene_data

        @contextmanager
        def _connect(_url):
            w.events.append("connect")
            yield _Conn(w.sql)

        def _tripwire(*_a, **_k):
            raise AssertionError("db._open must not be reached (db.connect is faked)")

        def _finish(conn, run_id, **kw):
            w.events.append("finish_sync_run")
            w.finish_calls.append((run_id, kw))

        monkeypatch.setattr(ingest, "InvestecClient", _FakeClient)
        monkeypatch.setattr(db, "_open", _tripwire)
        monkeypatch.setattr(db, "connect", _connect)
        monkeypatch.setattr(db, "init_db", lambda conn: None)
        monkeypatch.setattr(db, "start_sync_run", lambda conn, f, t: 1)
        monkeypatch.setattr(db, "get_account_ids", lambda conn: {"acc-1"})
        monkeypatch.setattr(db, "upsert_account", lambda conn, a: None)
        monkeypatch.setattr(db, "upsert_balance", lambda conn, a, b: None)
        monkeypatch.setattr(db, "upsert_transactions", lambda conn, a, items: len(items))
        monkeypatch.setattr(db, "finish_sync_run", _finish)
        settings = _settings(enabled)
        monkeypatch.setattr(Settings, "load", classmethod(lambda cls: settings))
        return w, settings
    return make


ARGS = argparse.Namespace(days=None, from_date="2026-01-01", to_date="2026-01-07",
                          resume=False, full=False)

SCENARIOS = {
    "off": dict(enabled=False),
    "on": dict(enabled=True),
    "on_raises": dict(enabled=True, bene_raises=RuntimeError("api down 9999 0000 1234")),
    "on_malformed": dict(enabled=True, bene_data={"bad": 1}),
    "on_none": dict(enabled=True, bene_data="junk"),
}


def _run_scenario(world_factory, capsys, **kw):
    w, settings = world_factory(**kw)
    assert cli.cmd_ingest(ARGS) == 0
    out = capsys.readouterr()
    summary = ingest.run_ingest(settings, from_date=dt.date(2026, 1, 1),
                                to_date=dt.date(2026, 1, 7))
    return w, out, summary


def test_ingest_output_frozen_off_vs_on(world_factory, capsys, caplog):
    caplog.set_level(logging.INFO)
    results = {name: _run_scenario(world_factory, capsys, **kw)
               for name, kw in SCENARIOS.items()}
    off_w, off_out, off_summary = results["off"]
    assert off_summary["new_accounts"] == []
    assert list(off_summary) == SUMMARY_KEYS
    for name, (w, out, summary) in results.items():
        assert out.out == off_out.out, name
        assert out.err == off_out.err, name
        assert summary == off_summary and list(summary) == SUMMARY_KEYS, name
        assert w.finish_calls == off_w.finish_calls, name
        assert all(kw["status"] == "success" for _, kw in w.finish_calls), name
    assert off_out.out.startswith("Ingest: {'from_date': '2026-01-01'")
    assert "9999" not in caplog.text and "1234" not in caplog.text


def test_flag_off_no_beneficiary_calls_or_sql(world_factory, capsys):
    w, _out, _summary = _run_scenario(world_factory, capsys, enabled=False)
    assert w.bene_calls == 0
    assert not any("beneficiar" in s.lower() for s in w.sql)


def test_flag_on_hook_actually_runs(world_factory, capsys):
    w, _out, _summary = _run_scenario(world_factory, capsys, enabled=True)
    assert w.bene_calls == 2  # once per run (cmd_ingest + run_ingest)
    lows = [" ".join(s.lower().split()) for s in w.sql]
    assert any(s.startswith("insert into beneficiaries") for s in lows)
    assert any(s.startswith("insert into beneficiary_matches") for s in lows)
    # Runs only after sync_runs was finalised; API call precedes its DB work.
    first = w.events.index("get_beneficiaries")
    assert w.events.index("finish_sync_run") < first
    assert w.events[first + 1] == "connect"


@pytest.mark.parametrize("name", ["on_raises", "on_malformed", "on_none"])
def test_flag_on_fetch_failure_is_non_fatal(world_factory, capsys, caplog, name):
    caplog.set_level(logging.WARNING)
    w, out, _summary = _run_scenario(world_factory, capsys, **SCENARIOS[name])
    assert not any("beneficiar" in s.lower() for s in w.sql)
    assert "Beneficiary" in caplog.text
    assert "Beneficiary" not in out.out
    assert "api down" not in caplog.text  # exception message never logged


def test_settings_flag_defaults_off(monkeypatch):
    for k, v in {"INVESTEC_CLIENT_ID": "x", "INVESTEC_CLIENT_SECRET": "x",
                 "INVESTEC_API_KEY": "x", "DATABASE_URL": "postgresql://fake"}.items():
        monkeypatch.setenv(k, v)
    monkeypatch.delenv("BENEFICIARY_MATCHING_ENABLED", raising=False)
    monkeypatch.delenv("PAYMENTS_BENEFICIARIES_FROM_API", raising=False)
    assert Settings.load().beneficiary_matching_enabled is False
    monkeypatch.setenv("BENEFICIARY_MATCHING_ENABLED", "true")
    assert Settings.load().beneficiary_matching_enabled is True
    assert Settings.load().payments_beneficiaries_from_api is False  # separate flag
    assert _settings(False).beneficiary_matching_enabled is False


# ── F10: Excel outputs unchanged (golden digest generated at base SHA) ──────

def _excel_fixture_df() -> pd.DataFrame:
    """Fixed synthetic transactions (fake names / numbers only)."""
    return pd.DataFrame(
        {
            "effective_date": pd.to_datetime(
                ["2026-01-01", "2026-01-02", "2026-01-02", "2026-01-03",
                 "2026-01-04", "2026-01-05"]
            ),
            "account_number": ["TEST-ACCT-01"] * 6,
            "account_name": ["Test Cheque"] * 6,
            "type": ["DEBIT", "DEBIT", "CREDIT", "DEBIT", "DEBIT", "CREDIT"],
            "transaction_type": ["OnlineBankingPayments", "CardPurchases", "Deposits",
                                 "FasterPay", "OnlineBankingTransfers",
                                 "OnlineBankingTransfers"],
            "description": ["PAYMENT TO ACME TRADING", "Corner Shop", "Salary",
                            "Transfer to J Smithers", "Transfer to savings",
                            "Transfer from savings"],
            "amount": [-1500.0, -85.5, 20000.0, -250.0, -1000.0, 400.0],
            "running_balance": [8500.0, 8414.5, 28414.5, 28164.5, None, 27564.5],
            "category": ["Other", "Groceries", "Income", "Other", "Transfers",
                         "Transfers"],
            "flow_type": ["external_outflow", "external_outflow", "external_inflow",
                          "external_outflow", "internal_transfer",
                          "internal_transfer"],
            "day_seq": [0, 0, 1, 0, 0, 0],
        }
    )


def _cell_value(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    return value


def _workbook_digest(path: Path) -> dict:
    wb = load_workbook(path)
    out = {"sheets": wb.sheetnames, "content": {}}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and c.number_format == "General":
                    continue
                cells.append([c.coordinate, _cell_value(c.value), c.number_format])
        widths = {k: v.width for k, v in sorted(ws.column_dimensions.items())}
        out["content"][ws.title] = {"cells": cells, "widths": widths,
                                    "freeze_panes": ws.freeze_panes}
    return out


def _build_excel_digests(tmp: Path) -> dict:
    from invespend import report, statements

    df = _excel_fixture_df()
    rpt = report.write_workbook(report.build_spend_summary(df), tmp / "report.xlsx")
    stmt = statements.build_account_statement(df, opening_balance=10000.0)
    acct = statements.Account(account_id="acc-1", account_number="TEST-ACCT-01",
                              account_name="Test Cheque")
    st = statements.write_statement_workbook(
        stmt, acct, dt.date(2026, 1, 1), dt.date(2026, 1, 7), tmp / "statement.xlsx")
    return {
        "report_columns": list(report.COLUMNS),
        "statement_columns": list(statements.STATEMENT_COLUMNS),
        "report": _workbook_digest(rpt),
        "statement": _workbook_digest(st),
    }


def _golden() -> dict:
    g = json.loads((FIXTURES / "excel_golden.json").read_text())
    g.pop("generated_at_base_sha")
    return json.loads(json.dumps(g))


def test_excel_outputs_unchanged(tmp_path, world_factory, capsys):
    off = _build_excel_digests(tmp_path / "off")
    _run_scenario(world_factory, capsys, enabled=True)  # hook ran, modules imported
    on = _build_excel_digests(tmp_path / "on")
    off_j, on_j = json.loads(json.dumps(off)), json.loads(json.dumps(on))
    assert off_j == on_j
    assert on_j == _golden()


def test_report_and_statement_columns_frozen():
    from invespend import report, statements

    assert report.COLUMNS == ["effective_date", "account_number", "account_name", "type",
                              "transaction_type", "description", "amount", "category",
                              "flow_type"]
    assert statements.STATEMENT_COLUMNS == ["Date", "Description", "Category", "Debit",
                                            "Credit", "Balance"]
